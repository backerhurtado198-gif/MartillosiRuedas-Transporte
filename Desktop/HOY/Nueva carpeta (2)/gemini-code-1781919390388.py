import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# 1. Inicializar la presentación
prs = Presentation()
prs.slide_width = Inches(13.33)  # Formato Panorámico 16:9
prs.slide_height = Inches(7.5)

# Colores de la marca (Tema oscuro elegante con naranja BiteFlow)
FONDO = RGBColor(15, 15, 15)
NARANJA = RGBColor(255, 140, 0)
BLANCO = RGBColor(255, 255, 255)
GRIS = RGBColor(150, 150, 150)
ROJO = RGBColor(220, 53, 69)

def aplicar_fondo_oscuro(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = FONDO

def agregar_titulo(slide, texto, color=NARANJA, size=40):
    txBox = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.83), Inches(1))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = texto
    p.font.bold = True
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.name = 'Arial'
    return txBox

# ----------------------------------------------------
# DIAPOSITIVA 1: PORTADA
# ----------------------------------------------------
blank_layout = prs.slide_layouts[6]
slide1 = prs.slides.add_slide(blank_layout)
aplicar_fondo_oscuro(slide1)

txBox = slide1.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(3.5))
tf = txBox.text_frame
p1 = tf.paragraphs[0]
p1.text = "BiteFlow Móvil"
p1.font.bold = True
p1.font.size = Pt(64)
p1.font.color.rgb = NARANJA

p2 = tf.add_paragraph()
p2.text = "El sistema de control total para tu negocio en la palma de tu mano"
p2.font.size = Pt(24)
p2.font.color.rgb = BLANCO
p2.space_before = Pt(20)

p3 = tf.add_paragraph()
p3.text = "Comida Rápida • Abastos • Bodegas • Minimercados"
p3.font.size = Pt(18)
p3.font.color.rgb = GRIS
p3.space_before = Pt(15)


# ----------------------------------------------------
# DIAPOSITIVA 2: EL PROBLEMA
# ----------------------------------------------------
slide2 = prs.slides.add_slide(blank_layout)
aplicar_fondo_oscuro(slide2)
agregar_titulo(slide2, "¿Cuánto dinero estás perdiendo hoy en descontrol?")

txBox = slide2.shapes.add_textbox(Inches(1), Inches(2.0), Inches(11.33), Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True

puntos_problema = [
    ("❌ Pérdidas ocultas (Merma):", " Panes duros, salchichas vencidas, víveres rotos o golpeados que se botan sin registrar."),
    ("❌ Cuentas en cuadernos o flojera en caja:", " Errores al dar el vuelto en Bs. o dólares, descuadres al final de la noche y falta de claridad."),
    ("❌ Quedarse sin mercancía:", " No saber el stock exacto provoca quedarse sin insumos clave un sábado por la noche o en hora pico.")
]

for titulo, desc in puntos_problema:
    p = tf.add_paragraph()
    p.space_before = Pt(20)
    run1 = p.add_run()
    run1.text = titulo
    run1.font.bold = True
    run1.font.size = Pt(22)
    run1.font.color.rgb = ROJO
    
    run2 = p.add_run()
    run2.text = desc
    run2.font.size = Pt(20)
    run2.font.color.rgb = BLANCO

# ----------------------------------------------------
# DIAPOSITIVA 3: LA SOLUCIÓN
# ----------------------------------------------------
slide3 = prs.slides.add_slide(blank_layout)
aplicar_fondo_oscuro(slide3)
agregar_titulo(slide3, "La Solución: BiteFlow Móvil")

txBox = slide3.shapes.add_textbox(Inches(1), Inches(2.0), Inches(11.33), Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True

puntos_solucion = [
    ("📱 100% Móvil y Ligera:", " Diseñada para trabajar rápido directo en tu teléfono Android, sin necesidad de costosas computadoras."),
    ("👥 Doble Rol de Seguridad:", " Interfaz simplificada para el Cajero (ventas rápidas) y un panel privado con clave para el Administrador."),
    ("⚙️ Multinegocio Adaptable:", " Ideal tanto para desbancar insumos de comida rápida como para llevar el conteo de piezas y víveres en abastos.")
]

for titulo, desc in puntos_solucion:
    p = tf.add_paragraph()
    p.space_before = Pt(20)
    run1 = p.add_run()
    run1.text = titulo
    run1.font.bold = True
    run1.font.size = Pt(22)
    run1.font.color.rgb = NARANJA
    
    run2 = p.add_run()
    run2.text = desc
    run2.font.size = Pt(20)
    run2.font.color.rgb = BLANCO


# ----------------------------------------------------
# DIAPOSITIVA 4: MÓDULO DE VENTAS
# ----------------------------------------------------
slide4 = prs.slides.add_slide(blank_layout)
aplicar_fondo_oscuro(slide4)
agregar_titulo(slide4, "🛒 Módulo de Ventas y Facturación Inteligente")

txBox = slide4.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True

puntos_ventas = [
    "• Registro de pedidos con un solo toque en la pantalla táctil.",
    "• Cálculo automático del total en moneda local (Bs.) para agilizar el vuelto.",
    "• Alerta de Stock Seguro: El sistema bloquea de inmediato la venta si detecta que no quedan ingredientes o insumos suficientes en el inventario."
]

for punto in puntos_ventas:
    p = tf.add_paragraph()
    p.text = punto
    p.font.size = Pt(22)
    p.font.color.rgb = BLANCO
    p.space_before = Pt(25)


# ----------------------------------------------------
# DIAPOSITIVA 5: INVENTARIO Y ANEXOS
# ----------------------------------------------------
slide5 = prs.slides.add_slide(blank_layout)
aplicar_fondo_oscuro(slide5)
agregar_titulo(slide5, "📦 Inventario Real y Botón de Anexar (+)")

txBox = slide5.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True

puntos_inv = [
    "• Control en caliente de insumos y productos terminados (Panes, salchichas, harinas, refrescos).",
    "• Alertas Visuales en Rojo cuando un producto tiene existencia crítica (menos de 10 unidades).",
    "• Botón de Anexar (+): Permite al dueño registrar nuevos víveres, marcas o mercancía que va llegando al local al instante sin enredos técnicos."
]

for punto in puntos_inv:
    p = tf.add_paragraph()
    p.text = punto
    p.font.size = Pt(22)
    p.font.color.rgb = BLANCO
    p.space_before = Pt(25)


# ----------------------------------------------------
# DIAPOSITIVA 6: REPORTE DE DAÑOS / MERMA
# ----------------------------------------------------
slide6 = prs.slides.add_slide(blank_layout)
aplicar_fondo_oscuro(slide6)
agregar_titulo(slide6, "🗑️ Control de Mermas y Productos Dañados", color=ROJO)

txBox = slide6.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True

puntos_merma = [
    "• Botón Exclusivo de Pérdidas (🗑️) al lado de cada producto del inventario.",
    "• ¿Se venció un empaque, se golpeó un vívere o se dañó el pan? Digitas la cantidad dañada y se descuenta automáticamente del stock.",
    "• Reporte de Auditoría: Se genera una sección en el historial para registrar cuándo y cuánto se perdió, evitando los descuadres fantasmas a fin de mes."
]

for punto in puntos_merma:
    p = tf.add_paragraph()
    p.text = punto
    p.font.size = Pt(22)
    p.font.color.rgb = BLANCO
    p.space_before = Pt(25)


# ----------------------------------------------------
# DIAPOSITIVA 7: CIERRE
# ----------------------------------------------------
slide7 = prs.slides.add_slide(blank_layout)
aplicar_fondo_oscuro(slide7)

txBox = slide7.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(4))
tf = txBox.text_frame
p1 = tf.paragraphs[0]
p1.text = "¡Lleva tu negocio al siguiente nivel!"
p1.font.bold = True
p1.font.size = Pt(44)
p1.font.color.rgb = NARANJA
p1.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = "Protege tus ganancias, elimina las pérdidas ocultas y vende más rápido."
p2.font.size = Pt(22)
p2.font.color.rgb = BLANCO
p2.space_before = Pt(20)
p2.alignment = PP_ALIGN.CENTER

p3 = tf.add_paragraph()
p3.text = "Agenda una demostración de 5 minutos en tu local hoy"
p3.font.bold = True
p3.font.size = Pt(20)
p3.font.color.rgb = GRIS
p3.space_before = Pt(30)
p3.alignment = PP_ALIGN.CENTER

# Guardar la presentación
prs.save("BiteFlow_Presentacion.pptx")
print("¡Listo, hermano! Diapositiva generada exitosamente como 'BiteFlow_Presentacion.pptx'")